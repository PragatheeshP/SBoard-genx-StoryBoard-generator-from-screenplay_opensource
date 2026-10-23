"""
Export module.
PDF export (one panel per page via ReportLab) and
PowerPoint export (four panels per slide via python-pptx).
"""
import logging
import uuid
from io import BytesIO
from pathlib import Path
from typing import List, Dict, Any

from PIL import Image
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image as RLImage,
    PageBreak,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from backend.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()


class ExportError(Exception):
    """Custom exception for export errors."""
    pass


# ─────────────────────────────────────────────────────────────────────────────
# PDF export
# ─────────────────────────────────────────────────────────────────────────────

class PDFExporter:
    """Exports a storyboard to PDF with one panel per page."""

    def __init__(self):
        """Initialise the PDF exporter."""
        self.output_dir = settings.output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        project_id: str,
        title: str,
        shots: List[Dict[str, Any]],
    ) -> str:
        """
        Export the storyboard to a PDF file.

        Each page contains: project title, shot number/metadata, and the
        generated panel image centred on the page.

        Args:
            project_id: Project identifier.
            title:       Title shown in the header of every page.
            shots:       Ordered list of shot dicts.

        Returns:
            Path to the exported PDF.
        """
        try:
            project_dir = self.output_dir / project_id
            project_dir.mkdir(parents=True, exist_ok=True)
            dest = project_dir / f"storyboard_{uuid.uuid4().hex[:8]}.pdf"

            doc = SimpleDocTemplate(
                str(dest),
                pagesize=A4,
                leftMargin=2 * cm,
                rightMargin=2 * cm,
                topMargin=2 * cm,
                bottomMargin=2 * cm,
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "title",
                parent=styles["Title"],
                fontSize=18,
                textColor=colors.HexColor("#1e1e32"),
                spaceAfter=6,
                alignment=TA_CENTER,
            )
            subtitle_style = ParagraphStyle(
                "subtitle",
                parent=styles["Normal"],
                fontSize=10,
                textColor=colors.HexColor("#5c5c7a"),
                spaceAfter=4,
                alignment=TA_CENTER,
            )
            caption_style = ParagraphStyle(
                "caption",
                parent=styles["Normal"],
                fontSize=9,
                textColor=colors.HexColor("#5c5c7a"),
                spaceBefore=4,
                alignment=TA_LEFT,
            )

            story = []
            page_w, _ = A4
            # Usable width in points (A4 = 595 pt, margins = 2*2cm = ~113 pt)
            usable_w = page_w - 4 * cm

            for idx, shot in enumerate(shots):
                # ── Header ────────────────────────────────────────────────
                story.append(Paragraph(title, title_style))
                story.append(
                    Paragraph(
                        f"Shot {shot.get('shot_number', idx + 1)} of {len(shots)}",
                        subtitle_style,
                    )
                )
                story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#c0c0d0")))
                story.append(Spacer(1, 0.3 * cm))

                # ── Panel image ──────────────────────────────────────────
                img_path = shot.get("image_path")
                if img_path and Path(img_path).exists():
                    rl_img = RLImage(img_path, width=usable_w, height=usable_w)
                    story.append(rl_img)
                else:
                    story.append(
                        Paragraph("[Panel image not yet generated]", caption_style)
                    )

                story.append(Spacer(1, 0.4 * cm))

                # ── Metadata captions ─────────────────────────────────────
                meta = [
                    f"<b>Type:</b> {shot.get('shot_type', '—').title()}",
                    f"<b>Angle:</b> {shot.get('camera_angle', '—').replace('-', ' ').title()}",
                    f"<b>Characters:</b> {', '.join(shot.get('characters', [])) or '—'}",
                    f"<b>Mood:</b> {shot.get('mood', '—').title()}",
                ]
                story.append(Paragraph("   |   ".join(meta), caption_style))
                story.append(Spacer(1, 0.3 * cm))

                action = shot.get("action", "")
                if action:
                    story.append(Paragraph(f"<i>{action}</i>", caption_style))

                # ── Page break after every panel (except last) ────────────
                if idx < len(shots) - 1:
                    story.append(PageBreak())

            doc.build(story)
            logger.info(f"PDF exported: {dest}")
            return str(dest)

        except Exception as e:
            logger.error(f"PDF export failed: {e}")
            raise ExportError(f"Failed to export PDF: {e}")


# ─────────────────────────────────────────────────────────────────────────────
# PowerPoint export
# ─────────────────────────────────────────────────────────────────────────────

class PPTXExporter:
    """Exports a storyboard to PowerPoint with four panels per slide."""

    PANELS_PER_SLIDE = 4

    # Slide layout in inches (16:9 widescreen)
    SLIDE_W = 13.33
    SLIDE_H = 7.5

    # Grid: 2 columns × 2 rows
    COLS = 2
    ROWS = 2
    PANEL_W = 5.5
    PANEL_H = 3.0
    H_PAD = 0.3
    V_PAD = 0.5
    TOP_PAD = 1.2   # space for slide title

    def __init__(self):
        """Initialise the PPTX exporter."""
        self.output_dir = settings.output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        project_id: str,
        title: str,
        shots: List[Dict[str, Any]],
    ) -> str:
        """
        Export the storyboard to a PowerPoint file.

        Every slide has a title bar and a 2×2 grid of panels with metadata.

        Args:
            project_id: Project identifier.
            title:       Presentation title shown on the first slide and in slide titles.
            shots:       Ordered list of shot dicts.

        Returns:
            Path to the exported PPTX.
        """
        try:
            prs = Presentation()
            prs.slide_width = Inches(self.SLIDE_W)
            prs.slide_height = Inches(self.SLIDE_H)

            blank_layout = prs.slide_layouts[6]  # blank layout

            # ── Title slide ──────────────────────────────────────────────
            self._add_title_slide(prs, blank_layout, title)

            # ── Content slides (4 panels each) ───────────────────────────
            for slide_idx in range(0, len(shots), self.PANELS_PER_SLIDE):
                panel_group = shots[slide_idx : slide_idx + self.PANELS_PER_SLIDE]
                slide_num = slide_idx // self.PANELS_PER_SLIDE + 1
                self._add_panel_slide(
                    prs, blank_layout, title, slide_num, panel_group
                )

            # ── Save ─────────────────────────────────────────────────────
            project_dir = self.output_dir / project_id
            project_dir.mkdir(parents=True, exist_ok=True)
            dest = project_dir / f"storyboard_{uuid.uuid4().hex[:8]}.pptx"
            prs.save(str(dest))
            logger.info(f"PPTX exported: {dest}")
            return str(dest)

        except Exception as e:
            logger.error(f"PPTX export failed: {e}")
            raise ExportError(f"Failed to export PowerPoint: {e}")

    # ─────────────────────── private helpers ─────────────────────────────────

    def _add_title_slide(self, prs, layout, title: str):
        """Add a cover / title slide."""
        slide = prs.slides.add_slide(layout)
        self._set_slide_bg(slide, "1e1e32")

        txb = slide.shapes.add_textbox(
            Inches(1), Inches(2.8), Inches(11.33), Inches(1.5)
        )
        tf = txb.text_frame
        tf.text = title
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(40)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xF0, 0xF0, 0xFF)

        sub = slide.shapes.add_textbox(
            Inches(1), Inches(4.2), Inches(11.33), Inches(0.6)
        )
        tf2 = sub.text_frame
        tf2.text = "Storyboard"
        tf2.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf2.paragraphs[0].runs[0].font.size = Pt(18)
        tf2.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xB0, 0xB0, 0xD0)

    def _add_panel_slide(
        self, prs, layout, title: str, slide_num: int, shots: List[Dict[str, Any]]
    ):
        """Add a content slide with up to 4 panels."""
        slide = prs.slides.add_slide(layout)
        self._set_slide_bg(slide, "f5f5fa")

        # Slide title bar
        title_box = slide.shapes.add_textbox(
            Inches(0.2), Inches(0.1), Inches(12.5), Inches(0.7)
        )
        tf = title_box.text_frame
        tf.text = f"{title}  ·  Slide {slide_num}"
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x32)

        for panel_idx, shot in enumerate(shots):
            row = panel_idx // self.COLS
            col = panel_idx % self.COLS

            left = Inches(self.H_PAD + col * (self.PANEL_W + self.H_PAD))
            top = Inches(
                self.TOP_PAD + row * (self.PANEL_H + self.V_PAD + 0.5)
            )  # +0.5 for caption row

            # Panel image
            img_path = shot.get("image_path")
            if img_path and Path(img_path).exists():
                slide.shapes.add_picture(
                    img_path,
                    left,
                    top,
                    Inches(self.PANEL_W),
                    Inches(self.PANEL_H),
                )
            else:
                # Placeholder box
                ph = slide.shapes.add_textbox(
                    left, top, Inches(self.PANEL_W), Inches(self.PANEL_H)
                )
                ph.text_frame.text = "Image pending"

            # Caption beneath the panel
            cap_top = top + Inches(self.PANEL_H + 0.05)
            cap_box = slide.shapes.add_textbox(
                left, cap_top, Inches(self.PANEL_W), Inches(0.4)
            )
            tf_cap = cap_box.text_frame
            shot_label = (
                f"SHOT {shot.get('shot_number', panel_idx + 1)}  ·  "
                f"{shot.get('shot_type', '').upper()}  ·  "
                f"{shot.get('camera_angle', '').replace('-', ' ').upper()}"
            )
            tf_cap.text = shot_label
            para = tf_cap.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x5C, 0x5C, 0x7A)

    @staticmethod
    def _set_slide_bg(slide, hex_colour: str):
        """Set the solid background colour of a slide."""
        fill = slide.background.fill
        fill.solid()
        r = int(hex_colour[0:2], 16)
        g = int(hex_colour[2:4], 16)
        b = int(hex_colour[4:6], 16)
        fill.fore_color.rgb = RGBColor(r, g, b)


# ─────────────────────────────────────────────────────────────────────────────
# Public service wrapper
# ─────────────────────────────────────────────────────────────────────────────

class ExportService:
    """High-level export service."""

    def __init__(self):
        """Initialise PDF and PPTX exporters."""
        self.pdf = PDFExporter()
        self.pptx = PPTXExporter()
        logger.info("ExportService initialised")

    def export_pdf(
        self, project_id: str, title: str, shots: List[Dict[str, Any]]
    ) -> str:
        """Export storyboard as PDF (one panel per page)."""
        return self.pdf.export(project_id, title, shots)

    def export_pptx(
        self, project_id: str, title: str, shots: List[Dict[str, Any]]
    ) -> str:
        """Export storyboard as PowerPoint (four panels per slide)."""
        return self.pptx.export(project_id, title, shots)


def get_export_service() -> ExportService:
    """Return a configured ExportService instance."""
    return ExportService()
